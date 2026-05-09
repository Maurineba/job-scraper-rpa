from utils.parse_salary import parse_salary

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from collections import Counter

from database.jobs_repository import get_all_jobs


def export_jobs(file_name="data/vagas.xlsx"):
   jobs = get_all_jobs()

   wb = Workbook()

   ws = wb.active

   ws.title = "DATA"

   ws.append(["Título", "Salário", "Modalidade", "Link"])

   for job in jobs:
      ws.append([
         job["titulo"],
         job["salario"],
         job["modalidade"],
         job["link"]
      ])

   dash = wb.create_sheet("DASHBOARD")

   total_vagas = len(jobs)

   salarios = [parse_salary(j["salario"]) for j in jobs]

   media_salario = sum(salarios) / len(salarios) if salarios else 0

   dash["A1"] = "Total de vagas"
   dash["B1"] = total_vagas
   dash["A2"] = "Média salarial"
   dash["B2"] = round(media_salario, 2)

   count_modalidade = Counter([j["modalidade"] for j in jobs])

   dash["A5"] = "Modalidade"
   dash["B5"] = "Quantidade"

   start_row = 6
   for i, (mod, qtd) in enumerate(count_modalidade.items(), start=start_row):
       dash[f"A{i}"] = mod
       dash[f"B{i}"] = qtd

   chart = PieChart()

   chart.title = "Vagas por Modalidade"

   data = Reference(
      dash,
      min_col=2,
      min_row=5,
      max_row=5 + len(count_modalidade)
   )

   labels = Reference(
      dash,
      min_col=1,
      min_row=6,
      max_row=5 + len(count_modalidade)
   )

   chart.add_data(data, titles_from_data=True)
   chart.set_categories(labels)
   dash.add_chart(chart, "D10")

   wb.save(file_name)

   return file_name

